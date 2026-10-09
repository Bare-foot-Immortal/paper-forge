# -*- coding: utf-8 -*-
"""Word(.docx) 与 PDF 题库导入测试（先于实现编写，用于驱动开发）。

分三层：
* 段落文本解析器单元测试（各种答案形态、章节切换、粘连行等）；
* 合成 Word 样例（``fixtures/``，由 tools/make_samples.py 生成）的集成测试；
* PDF 与表格型 Word 的导入测试。

运行：python -m unittest tests.test_word_pdf_import -v
"""
from __future__ import annotations

import logging
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

from paperforge.bank_io import (  # noqa: E402
    SUPPORTED_SUFFIXES,
    TextLine,
    lines_to_questions,
    load_bank,
)
from paperforge.models import QType  # noqa: E402

logging.getLogger("pypdf").setLevel(logging.CRITICAL)   # 坏文件用例会触发 pypdf 告警，测试中静音

FIXTURES = ROOT / "fixtures"
SAMPLES = ROOT / "samples"

WORD_SAMPLES = {
    "a": FIXTURES / "样例A_段落文本型.docx",
    "b": FIXTURES / "样例B_选项同行.docx",
    "c": FIXTURES / "样例C_题干答案分行.docx",
    "d": FIXTURES / "样例D_选项无分隔符.docx",
    "e": FIXTURES / "样例E_表格型.docx",
}
PDF_SAMPLE = FIXTURES / "样例题库.pdf"
BANK_XLSX = FIXTURES / "样例题库.xlsx"


def _ensure_fixtures() -> None:
    """夹具缺失时立即生成（必须在类定义/装饰器求值之前调用）。"""
    if BANK_XLSX.exists():
        return
    import subprocess

    script = ROOT / "tools" / "make_samples.py"
    if script.exists():
        subprocess.run([sys.executable, str(script), "--fixtures-only"],
                       cwd=str(ROOT), check=False, capture_output=True)


_ensure_fixtures()




def lines(*texts: str) -> list[TextLine]:
    return [TextLine(text=t) for t in texts]


def parse(texts: list[str]):
    return lines_to_questions(lines(*texts))


# ================================================================ 段落解析器
class TestTextParser(unittest.TestCase):
    """段落文本题库解析（Word 段落型 / PDF 文本型共用）。"""

    def test_embedded_single_answer(self):
        qs, issues, _ = parse([
            "一、单项选择题",
            "1、深基坑工程中，开挖深度超过（A）米（含）的基坑支护工程属于危大工程？",
            "A.3", "B.5", "C.7", "D.10",
        ])
        self.assertEqual(issues, [])
        self.assertEqual(len(qs), 1)
        q = qs[0]
        self.assertIs(q.qtype, QType.SINGLE)
        self.assertEqual(q.answer_letters, ["A"])
        self.assertEqual(len(q.options), 4)

    def test_embedded_multi_answer(self):
        qs, issues, _ = parse([
            "二、多项选择题",
            "1、下列属于深基坑工程危大工程的有（ABCD）？",
            "A.开挖深度超过 3 米的基坑支护", "B.开挖深度超过 5 米的基坑支护",
            "C.地下暗挖工程", "D.顶管工程",
        ])
        self.assertEqual(issues, [])
        self.assertEqual(len(qs), 1)
        self.assertIs(qs[0].qtype, QType.MULTIPLE)
        self.assertEqual(qs[0].answer_letters, ["A", "B", "C", "D"])

    def test_embedded_answer_with_spaces(self):
        """无分隔符选项样例的形态：答案写在带空格的全角括号里，选项到 E。"""
        qs, issues, _ = parse([
            "二、多项选择题练习",
            "1.根据管理规定，输变电工程建设期间（  BCE ）应按规定配备专职质量管理人员。",
            "A．甲单位    B. 乙单位    C.班组长    D.项目经理    E.专业技术负责人",
        ])
        self.assertEqual(issues, [])
        self.assertEqual(len(qs), 1)
        self.assertEqual(qs[0].answer_letters, ["B", "C", "E"])
        self.assertEqual(len(qs[0].options), 5)

    def test_answer_stripped_from_stem(self):
        """题干里的答案必须抹掉（否则卷面直接剧透）。"""
        qs, _, _ = parse([
            "一、单项选择题",
            "1、这是施工管理阶段的（C）。",
            "A.甲", "B.乙", "C.丙", "D.丁",
        ])
        self.assertNotIn("（C）", qs[0].stem)
        self.assertNotIn("C）", qs[0].stem)
        self.assertIn("（", qs[0].stem)          # 保留作答空位
        self.assertEqual(qs[0].answer_letters, ["C"])

    def test_answer_line_style(self):
        """「题干：/选项/答案：」版式：三段各自成行。"""
        qs, issues, _ = parse([
            "第一部分 单项选择题",
            "题干：模板安装作业开始前应当首先完成的工作是？",
            "A.先行安排材料验收 B.完成方案审批与安全技术交底 C.直接组织班组作业 D.等待口头同意",
            "答案：B",
            "来源：《合成示例管理规定》-第一章第一条",
            "题干：混凝土浇筑前应当确认的事项是？",
            "A.模板与钢筋已验收合格 B.材料价格已谈定 C.车辆数量已确定 D.天气已放晴",
            "答案：A",
        ])
        self.assertEqual(issues, [])
        self.assertEqual(len(qs), 2)
        self.assertEqual(qs[0].answer_letters, ["B"])
        self.assertEqual(qs[1].answer_letters, ["A"])
        self.assertIs(qs[0].qtype, QType.SINGLE)
        self.assertIn("合成示例管理规定", qs[0].analysis)

    def test_answer_line_with_fixed_marker(self):
        """`答案：BC（定）` 的「定」标记必须保留为固定顺序题。"""
        qs, issues, _ = parse([
            "二、多项选择题",
            "1、构支架吊装时，应从柱四周向杯口放入（ ）个木楔，并使立柱基本（ ）后，才能松大钩。",
            "A.3～4", "B.4～5", "C.垂直", "D.平行",
            "答案：BC（定）",
        ])
        self.assertEqual(issues, [])
        self.assertEqual(len(qs), 1)
        self.assertTrue(qs[0].fixed_order)
        self.assertEqual(qs[0].answer_letters, ["B", "C"])

    def test_judge_with_symbols(self):
        qs, issues, _ = parse([
            "三、判断题",
            "1、未经论证的情况下压缩合同约定工期，属于重大事故隐患。（√）",
            "2、可以在电缆沟内充装易燃易爆危险品。（×）",
            "3、施工单位应当对危大工程进行监测和安全巡视。（正确）",
        ])
        self.assertEqual(issues, [])
        self.assertEqual(len(qs), 3)
        for q in qs:
            self.assertIs(q.qtype, QType.JUDGE)
        self.assertTrue(qs[0].judge_answer)
        self.assertFalse(qs[1].judge_answer)
        self.assertTrue(qs[2].judge_answer)

    def test_empty_bracket_is_not_answer(self):
        """空括号 `（  ）` 是作答位、不是答案 → 应报“未找到答案”。"""
        qs, issues, _ = parse([
            "一、单项选择题",
            "1、这是一道没有给出答案的题目（  ）。",
            "A.甲", "B.乙", "C.丙", "D.丁",
        ])
        self.assertEqual(qs, [])
        self.assertEqual(len(issues), 1)
        self.assertIn("答案", str(issues[0]))

    def test_section_title_controls_type(self):
        qs, _, _ = parse([
            "第二部分 多项选择题",
            "题干：以下哪些属于安全管理类隐患？",
            "A.甲 B.乙 C.丙 D.丁",
            "答案：AB",
        ])
        self.assertIs(qs[0].qtype, QType.MULTIPLE)

    def test_single_section_with_multi_answer_is_tolerated(self):
        """章节写“单选”但答案是多字母 → 容错为多选，避免丢题。"""
        qs, issues, meta = parse([
            "一、单项选择题",
            "1、以下哪些属于危大工程（ABCD）？",
            "A.甲", "B.乙", "C.丙", "D.丁",
        ])
        self.assertEqual(len(qs), 1)
        self.assertIs(qs[0].qtype, QType.MULTIPLE)

    def test_skip_subjective_section(self):
        """主观题/案例题不属于三种客观题型，应跳过。"""
        qs, issues, meta = parse([
            "一、单项选择题",
            "1、客观题（A）。", "A.甲", "B.乙",
            "四、主观题：",
            "1、请简述危大工程的管理流程。",
            "2、请分析上述案例的责任划分。",
        ])
        self.assertEqual(len(qs), 1)
        self.assertIs(qs[0].qtype, QType.SINGLE)

    def test_inline_options_split(self):
        qs, _, _ = parse([
            "一、单项选择题",
            "1、题干（B）。",
            "A.甲方案说明    B.乙方案说明    C.丙方案说明    D.丁方案说明",
        ])
        self.assertEqual(qs[0].options,
                         ["甲方案说明", "乙方案说明", "丙方案说明", "丁方案说明"])
        self.assertEqual(qs[0].answer_letters, ["B"])

    def test_per_line_options(self):
        qs, _, _ = parse([
            "一、单项选择题",
            "1、题干（C）。", "A.甲", "B.乙", "C.丙", "D.丁",
        ])
        self.assertEqual(qs[0].options, ["甲", "乙", "丙", "丁"])

    def test_glued_next_question_is_split(self):
        """选项行尾部粘连下一题题干时必须切开。"""
        qs, issues, _ = parse([
            "二、多项选择题",
            "1.专项施工方案实施前，（  AB  ）应当向现场管理人员进行方案交底。",
            "A．编制人员    B.项目技术负责人    C.班组长    D.项目经理",
            "2.《大体积混凝土施工规范》中规定水泥进场时应检查哪些内容（ ABCD ）。",
            "A.品种    B.强度等级    C.出厂日期    D.包装编号",
        ])
        self.assertEqual(len(qs), 2)
        self.assertEqual(qs[0].answer_letters, ["A", "B"])
        self.assertEqual(qs[1].answer_letters, ["A", "B", "C", "D"])
        self.assertIn("大体积混凝土", qs[1].stem)
        self.assertEqual(len(qs[0].options), 4)
        self.assertNotIn("大体积", "".join(qs[0].options))

    def test_multiline_stem(self):
        qs, _, _ = parse([
            "一、单项选择题",
            "1、第一行题干内容，",
            "第二行继续描述（D）。",
            "A.甲", "B.乙", "C.丙", "D.丁",
        ])
        self.assertEqual(len(qs), 1)
        self.assertIn("第二行继续描述", qs[0].stem)
        self.assertEqual(qs[0].answer_letters, ["D"])

    def test_title_and_blank_lines_ignored(self):
        qs, issues, _ = parse([
            "某单位题库标题",
            "",
            "一、单项选择题",
            "1、题干（A）。", "A.甲", "B.乙",
        ])
        self.assertEqual(len(qs), 1)
        self.assertEqual(issues, [])

    def test_question_number_prefix_stripped(self):
        qs, _, _ = parse(["一、单项选择题", "1、题干内容（A）。", "A.甲", "B.乙"])
        self.assertFalse(qs[0].stem.startswith("1"))
        self.assertTrue(qs[0].stem.startswith("题干内容"))

    def test_bad_row_reported_with_line_number(self):
        qs, issues, _ = parse([
            "一、单项选择题",
            "1、正常题（A）。", "A.甲", "B.乙",
            "2、这道题没有答案（  ）。", "A.甲", "B.乙",
        ])
        self.assertEqual(len(qs), 1)
        self.assertEqual(len(issues), 1)
        self.assertGreater(issues[0].row, 1)

    def test_option_letters_and_answers_sorted(self):
        qs, _, _ = parse([
            "二、多项选择题",
            "1、题干（D、B）？", "A.甲", "B.乙", "C.丙", "D.丁",
        ])
        self.assertEqual(qs[0].answer_letters, ["B", "D"])


# ================================================================ Word 样例（合成夹具）
class TestWordSamples(unittest.TestCase):
    """用 tools/make_samples.py 生成的合成 Word 样例做集成测试（可安全公开）。"""

    EXPECTED = {"a": 8, "b": 3, "c": 4, "d": 3, "e": 4}

    @classmethod
    def setUpClass(cls):
        cls.results = {}
        for key, path in WORD_SAMPLES.items():
            if path.exists():
                cls.results[key] = load_bank(path)
        if len(cls.results) < len(cls.EXPECTED):
            raise unittest.SkipTest("未找到 Word 夹具，请先运行 tools/make_samples.py")

    def _need(self, key: str):
        if key not in self.results:
            self.skipTest(f"样例不存在：{WORD_SAMPLES[key]}")
        return self.results[key]

    def test_all_samples_parse(self):
        for key, res in self.results.items():
            with self.subTest(sample=key):
                self.assertEqual(res.issues, [], f"{key} 不应产生坏行：{res.issues[:2]}")
                self.assertGreaterEqual(len(res.questions), 3, f"{key} 解析出的题目过少")

    def test_expected_counts(self):
        for key, n in self.EXPECTED.items():
            with self.subTest(sample=key):
                self.assertEqual(len(self._need(key).questions), n)

    def test_all_questions_valid(self):
        for key, res in self.results.items():
            for q in res.questions:
                with self.subTest(sample=key, stem=q.stem[:20]):
                    self.assertEqual(q.validate(), [])

    def test_answers_not_leaked_into_stem(self):
        """题干里不应残留“（X）”形式的答案。"""
        for key, res in self.results.items():
            leaked = [q for q in res.questions
                      if q.qtype is not QType.JUDGE
                      and f"（{''.join(q.answer_letters)}）" in q.stem]
            with self.subTest(sample=key):
                self.assertEqual(leaked, [], f"{key} 有 {len(leaked)} 道题的题干里残留答案")

    def test_sample_a_has_three_types(self):
        types = {q.qtype for q in self._need("a").questions}
        self.assertEqual(types, {QType.SINGLE, QType.MULTIPLE, QType.JUDGE})

    def test_sample_a_embedded_answers(self):
        res = self._need("a")
        hit = [q for q in res.questions if "钢筋绑扎作业开始前" in q.stem]
        self.assertTrue(hit, "应能解析出嵌在题干括号里的答案")
        self.assertIs(hit[0].qtype, QType.SINGLE)
        self.assertEqual(hit[0].answer_letters, ["A"])
        multi = [q for q in res.questions if "安全检查应当重点关注" in q.stem]
        self.assertTrue(multi)
        self.assertEqual(multi[0].answer_letters, ["A", "B", "C"])

    def test_sample_c_answer_line_and_source(self):
        res = self._need("c")
        hit = [q for q in res.questions if "钢筋绑扎作业开始前" in q.stem]
        self.assertTrue(hit, "应能解析「题干：/答案：」版式")
        self.assertEqual(hit[0].answer_letters, ["A"])
        self.assertIn("合成示例管理规定", hit[0].analysis, "「来源：」行应作为解析保留")

    def test_sample_d_wide_answer_and_glue(self):
        res = self._need("d")
        self.assertEqual([q.answer_letters for q in res.questions],
                         [["A", "B", "C"], ["A", "B", "D"], ["B", "C", "E"]])
        self.assertEqual(len(res.questions[2].options), 5, "第 3 题应有 A～E 五个选项")

    def test_sample_e_table_with_combined_options(self):
        res = self._need("e")
        self.assertEqual([q.qtype for q in res.questions],
                         [QType.SINGLE, QType.SINGLE, QType.SINGLE, QType.MULTIPLE])
        self.assertTrue(res.questions[3].fixed_order, "表格里的「（定）」标记必须保留")
        self.assertEqual(res.questions[3].answer_letters, ["A", "B", "C"])


# ================================================================ 表格型 Word
class TestDocxTable(unittest.TestCase):
    """表格型 Word 题库（含“选项合并在一列”的宽容处理）。"""

    def _make_docx(self, path: Path, header: list[str], rows: list[list[str]]) -> None:
        from docx import Document

        doc = Document()
        table = doc.add_table(rows=1, cols=len(header))
        for i, text in enumerate(header):
            table.cell(0, i).text = text
        for row in rows:
            cells = table.add_row().cells
            for i, text in enumerate(row):
                cells[i].text = text
        doc.save(str(path))

    def test_table_style(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "表格题库.docx"
            self._make_docx(path,
                            ["题型", "题目标题", "选项A", "选项B", "选项C", "选项D", "解析", "答案"],
                            [["单选题", "题目一", "甲", "乙", "丙", "丁", "解析一", "B"],
                             ["多选题", "题目二", "甲", "乙", "丙", "丁", "解析二", "AC（定）"],
                             ["判断题", "题目三", "", "", "", "", "解析三", "正确"]])
            res = load_bank(path)
            self.assertEqual(len(res.questions), 3)
            self.assertEqual([q.qtype for q in res.questions],
                             [QType.SINGLE, QType.MULTIPLE, QType.JUDGE])
            self.assertEqual(res.questions[1].answer_letters, ["A", "C"])
            self.assertTrue(res.questions[1].fixed_order)

    def test_table_with_combined_option_column(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "合并选项列.docx"
            self._make_docx(path, ["题型", "题目标题", "选项", "答案"],
                            [["单选题", "题目一", "A.甲 B.乙 C.丙 D.丁", "C"]])
            res = load_bank(path)
            self.assertEqual(len(res.questions), 1)
            self.assertEqual(len(res.questions[0].options), 4)
            self.assertEqual(res.questions[0].answer_letters, ["C"])


# ================================================================ PDF
class TestPdfImport(unittest.TestCase):
    @unittest.skipUnless(PDF_SAMPLE.exists(), f"未找到 PDF 样例：{PDF_SAMPLE}")
    def test_builtin_pdf(self):
        res = load_bank(PDF_SAMPLE)
        self.assertGreaterEqual(len(res.questions), 8)
        types = {q.qtype for q in res.questions}
        self.assertEqual(types, {QType.SINGLE, QType.MULTIPLE, QType.JUDGE})
        fixed = [q for q in res.questions if q.fixed_order]
        self.assertTrue(fixed, "PDF 中的「（定）」题必须保留固定顺序标记")
        self.assertEqual(fixed[0].answer_letters, ["B", "C"])

    def test_fixture_bank_import(self):
        """合成夹具题库（52 题）整体导入：三题型齐全、含 1 道固定顺序题。"""
        if not BANK_XLSX.exists():
            self.skipTest("未找到夹具题库")
        res = load_bank(BANK_XLSX)
        self.assertEqual(len(res.issues), 0)
        self.assertEqual(len(res.questions), 52)
        counts = {t: len([q for q in res.questions if q.qtype is t])
                  for t in (QType.SINGLE, QType.MULTIPLE, QType.JUDGE)}
        self.assertEqual(counts, {QType.SINGLE: 26, QType.MULTIPLE: 14, QType.JUDGE: 12})
        self.assertEqual(len([q for q in res.questions if q.fixed_order]), 1)

    def test_broken_pdf_gives_readable_error(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "坏文件.pdf"
            path.write_bytes(b"this is not a pdf at all")
            with self.assertRaises(ValueError) as ctx:
                load_bank(path)
            self.assertIn("PDF", str(ctx.exception))

    def test_scanned_pdf_hint(self):
        """无文本层的 PDF 应给出“可能是扫描件”的可读提示。"""
        from pypdf import PdfWriter

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "无文本.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=200, height=200)
            with path.open("wb") as fh:
                writer.write(fh)
            with self.assertRaises(ValueError) as ctx:
                load_bank(path)
            self.assertTrue(any(k in str(ctx.exception) for k in ("扫描", "未识别", "题目")))


# ================================================================ 格式注册
class TestSupportedFormats(unittest.TestCase):
    def test_new_formats_registered(self):
        for suffix in (".docx", ".docm", ".pdf"):
            self.assertIn(suffix, SUPPORTED_SUFFIXES)

    def test_unknown_format_message(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "题库.xyz"
            path.write_text("x", encoding="utf-8")
            with self.assertRaises(ValueError) as ctx:
                load_bank(path)
            self.assertIn(".docx", str(ctx.exception))
            self.assertIn(".pdf", str(ctx.exception))


if __name__ == "__main__":
    unittest.main(verbosity=2)

