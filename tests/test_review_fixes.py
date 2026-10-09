# -*- coding: utf-8 -*-
"""代码审查（Agent Team）发现问题的回归测试（抽题匠）。

解析层用例与刷题匠共用同一实现；导出层用例覆盖两个高危缺陷：
多轮导出文件名冲突导致整轮被覆盖、合并导出且不单独出答案时文档缺答案。
"""
from __future__ import annotations

import shutil
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import Workbook  # noqa: E402

from paperforge.bank_io import TextLine, lines_to_questions, load_bank  # noqa: E402
from paperforge.models import QType  # noqa: E402

FIXTURE_BANK = ROOT / "fixtures" / "样例题库.xlsx"


def xlsx(path: Path, rows: list[list]) -> Path:
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


HEADER12 = ["题型", "题目标题", "选项A", "选项B", "选项C", "选项D", "解析", "答案"]


def bank_rows(n: int, start: int = 0) -> list[list]:
    return [HEADER12] + [["单选题", f"题目{i}", "甲", "乙", "丙", "丁", "", "A"]
                         for i in range(start, start + n)]


# ================================================================ 解析层
class TestParserReviewFixes(unittest.TestCase):
    def test_letter_answer_without_section_is_single_choice(self):
        """无章节标题时，答案 `（B）` 不得把单选题误判为判断题（并丢掉选项）。"""
        qs, issues, _ = lines_to_questions([
            TextLine("1、模板安装作业开始前应当首先完成的工作是（B）。"),
            TextLine("A.先行安排材料验收    B.完成方案审批与安全技术交底"),
            TextLine("C.直接组织班组作业    D.等待口头同意")])
        self.assertEqual(issues, [])
        self.assertEqual(len(qs), 1)
        self.assertIs(qs[0].qtype, QType.SINGLE)
        self.assertEqual(len(qs[0].options), 4)
        self.assertEqual(qs[0].answer_letters, ["B"])

    def test_letter_answer_with_judge_like_options_stays_judge(self):
        """选项本身就是「正确/错误」时，`（A）` 仍应判为判断题（修复不得过头）。"""
        qs, issues, _ = lines_to_questions([
            TextLine("1、未经论证压缩合同约定工期属于重大事故隐患。（A）"),
            TextLine("A.正确    B.错误")])
        self.assertEqual(issues, [])
        self.assertEqual(len(qs), 1)
        self.assertIs(qs[0].qtype, QType.JUDGE)
        self.assertTrue(qs[0].judge_answer)

    def test_judge_symbols_still_work(self):
        qs, issues, _ = lines_to_questions([
            TextLine("三、判断题"),
            TextLine("1、可以在电缆沟内充装易燃易爆危险品。（×）"),
            TextLine("2、施工单位应当对危大工程进行监测。（√）")])
        self.assertEqual(issues, [])
        self.assertEqual([q.judge_answer for q in qs], [False, True])

    def test_headerless_odd_width_is_rejected_loudly(self):
        """无列头且列数不在已知模板内时：明确报错，不得按位置猜列。"""
        tmp = Path(tempfile.mkdtemp(prefix="pd_rev_"))
        self.addCleanup(shutil.rmtree, tmp, True)
        for width, rows in (
                (6, [["单选题", "题目一", "甲", "乙", "丙", "丁"]]),
                (7, [["单选题", "题目一", "甲", "乙", "丙", "丁", "A"]])):
            with self.subTest(width=width):
                p = xlsx(tmp / f"w{width}.xlsx", rows)
                with self.assertRaises(ValueError) as ctx:
                    load_bank(p)
                self.assertIn("列头", str(ctx.exception))

    def test_headerless_standard_layout_still_supported(self):
        """12 列标准布局（无列头）仍按位置回退解析，不受上面收紧影响。"""
        tmp = Path(tempfile.mkdtemp(prefix="pd_rev_"))
        self.addCleanup(shutil.rmtree, tmp, True)
        rows = [["单选题", "题目一", "甲", "乙", "丙", "丁", "", "", "", "", "解析", "A"],
                ["判断题", "题目二", "", "", "", "", "", "", "", "", "", "正确"]]
        res = load_bank(xlsx(tmp / "std.xlsx", rows))
        self.assertEqual(len(res.questions), 2)
        self.assertEqual(len(res.issues), 0)

    def test_combined_option_column_below_title_rows(self):
        """表头不在第 1 行（上方有标题/空行）时，「选项」合并列仍应展开。"""
        tmp = Path(tempfile.mkdtemp(prefix="pd_rev_"))
        self.addCleanup(shutil.rmtree, tmp, True)
        res = load_bank(xlsx(tmp / "combo.xlsx", [
            ["2026 年度安全知识题库"], [],
            ["题型", "题目标题", "选项", "答案"],
            ["单选题", "题目一", "A.甲 B.乙 C.丙 D.丁", "B"],
            ["多选题", "题目二", "A.甲 B.乙 C.丙 D.丁", "AC（定）"]]))
        self.assertEqual(len(res.issues), 0)
        self.assertEqual(len(res.questions), 2)
        self.assertEqual(res.questions[0].options, ["甲", "乙", "丙", "丁"])
        self.assertTrue(res.questions[1].fixed_order)

    def test_answer_line_variants(self):
        """`【答案】A` / `答案 A` / `[答案]A` 等写法都应识别。"""
        for text in ("【答案】A", "答案 A", "[答案]A", "答案：A", "正确答案：A"):
            with self.subTest(line=text):
                qs, issues, _ = lines_to_questions([
                    TextLine("一、单项选择题"), TextLine("1、题干一？"),
                    TextLine("A.甲 B.乙"), TextLine(text)])
                self.assertEqual(issues, [], f"{text} 未识别为答案行")
                self.assertEqual(qs[0].answer_letters, ["A"])

    def test_glue_splits_short_next_stem(self):
        """选项行尾部粘连"下一题题干很短（≤8 字）"时也必须切开。"""
        qs, issues, _ = lines_to_questions([
            TextLine("一、单项选择题"),
            TextLine("1、题干一（A）。"),
            TextLine("A.甲    B.乙    C.丙    D.丁5、正确（B）。"),
            TextLine("A.甲    B.乙")])
        self.assertEqual(len(qs), 2)
        self.assertEqual(qs[1].answer_letters, ["B"])
        self.assertIn("正确", qs[1].stem)


# ================================================================ 导出层
class TestExportReviewFixes(unittest.TestCase):
    """覆盖 Agent Team 审查发现的导出缺陷（数据丢失级）。"""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="pf_rev_"))
        self.addCleanup(shutil.rmtree, self.tmp, True)
        rows = [["题型", "题目标题", "选项A", "选项B", "选项C", "选项D", "解析", "答案"]]
        for i in range(30):
            rows.append(["单选题", f"题目{i}", "甲", "乙", "丙", "丁", "", "ABCD"[i % 4]])
        self.questions = load_bank(xlsx(self.tmp / "bank.xlsx", rows)).questions

    def _report(self, rounds: int, round_length: int = 2):
        from paperforge.generator import Blueprint, GenOptions, generate
        return generate(self.questions, GenOptions(blueprint=Blueprint(5, 0, 0), seed=1,
                                                   rounds=rounds, round_length=round_length))

    def test_multi_round_export_does_not_overwrite(self):
        """多轮分文件导出：报告的文件数必须等于磁盘上的文件数（不得互相覆盖）。"""
        from paperforge.exporter import ExportOptions, export_all
        report = self._report(rounds=3)
        out = self.tmp / "multi"
        res = export_all(report, ExportOptions(out_dir=out, formats=("txt", "html", "docx"),
                                               answers_separate=True))
        disk = [p for p in out.rglob("*") if p.is_file()]
        self.assertEqual(len(res.files), len(disk), "存在同名覆盖")
        self.assertEqual(len({p.name for p in disk}), len(disk), "文件名重复")
        self.assertEqual(len(res.errors), 0)

    def test_single_round_filenames_unchanged(self):
        """单轮导出仍沿用「A卷」式文件名（向后兼容）。"""
        from paperforge.exporter import ExportOptions, export_all
        report = self._report(rounds=1)
        out = self.tmp / "single"
        export_all(report, ExportOptions(out_dir=out, formats=("txt",), answers_separate=False))
        names = sorted(p.name for p in out.glob("*.txt"))
        self.assertTrue(all("卷" in n for n in names), names)
        self.assertFalse(any(n[:2].isdigit() for n in names), f"单轮不应带序号：{names}")

    def test_merge_docx_includes_answers_when_not_separate(self):
        """合并导出且不单独出答案文件时，docx 必须包含参考答案（原来完全缺失）。"""
        from paperforge.exporter import ExportOptions, export_all
        report = self._report(rounds=1)
        out = self.tmp / "merged"
        export_all(report, ExportOptions(out_dir=out, formats=("docx",), merge=True,
                                        answers_separate=False))
        files = list(out.glob("*.docx"))
        self.assertEqual(len(files), 1)
        from docx import Document
        text = "\n".join(p.text for p in Document(str(files[0])).paragraphs)
        self.assertIn("参考答案", text)

    def test_merge_docx_separate_answers_not_in_paper(self):
        """答案单独成文件时，合并的试卷 docx 里不应出现参考答案（避免误带答案）。"""
        from paperforge.exporter import ExportOptions, export_all
        report = self._report(rounds=1)
        out = self.tmp / "merged_sep"
        export_all(report, ExportOptions(out_dir=out, formats=("docx",), merge=True,
                                        answers_separate=True))
        from docx import Document
        paper_files = [p for p in out.glob("*.docx") if "答案" not in p.name]
        self.assertTrue(paper_files)
        text = "\n".join(p.text for p in Document(str(paper_files[0])).paragraphs)
        self.assertNotIn("参考答案", text)


if __name__ == "__main__":
    unittest.main(verbosity=2)
