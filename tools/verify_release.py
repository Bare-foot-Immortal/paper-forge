# -*- coding: utf-8 -*-
"""发布校验工具：回读 exe/程序生成的 Word 试卷，独立核验组卷是否正确。

用法（在已生成试卷后执行）::

    python tools/verify_release.py --bank fixtures/样例题库.xlsx --out ..\\_recon\\exe_check --seed 12345

核验项：
  A. 每张卷子题量 / 题型数量 == 蓝图
  B. 一轮所有卷子的题目并集 == 题库全部题目（覆盖性）
  C. 卷面选项顺序与答案卷字母自洽（打乱后答案重映射正确）
  D. 与同种子的源码生成结果逐题一致（发布一致性）
  E. 卷头 / 分值 / 判断题作答括号等排版要素

退出码 0 表示全部通过，1 表示存在失败项。
"""
from __future__ import annotations

import argparse
import re
import sys
import unicodedata
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from docx import Document  # noqa: E402

from paperforge.bank_io import load_bank  # noqa: E402
from paperforge.generator import GenOptions, generate  # noqa: E402
from paperforge.models import Blueprint, QType  # noqa: E402

failures: list[str] = []


def check(cond: bool, label: str) -> None:
    print(("  PASS  " if cond else "  FAIL  ") + label)
    if not cond:
        failures.append(label)


def norm(s: str) -> str:
    s = unicodedata.normalize("NFKC", s or "")
    return re.sub(r"\s+", "", s.replace("\u3000", " ").replace("\xa0", " ")).strip()


def strip_blank(stem: str) -> str:
    return re.sub(r"[（(]\s*[）)]\s*$", "", stem.strip()).strip()


def parse_paper(path: Path) -> list[dict]:
    paras = [p.text.strip() for p in Document(path).paragraphs]
    items: list[dict] = []
    cur: "dict | None" = None
    for text in paras:
        if re.match(r"^\d+\.\s", text):
            if cur:
                items.append(cur)
            cur = {"stem": re.sub(r"^\d+\.\s*", "", text), "options": []}
        elif re.match(r"^[A-H]\.\s", text) and cur is not None:
            for part in re.split(r"\s{2,}(?=[A-H]\.\s)", text):
                m = re.match(r"^([A-H])\.\s*(.*)$", part.strip())
                if m:
                    cur["options"].append((m.group(1), m.group(2).strip()))
    if cur:
        items.append(cur)
    return items


def parse_answers(path: Path) -> list[str]:
    out = []
    for p in Document(path).paragraphs:
        m = re.match(r"^(\d+)\.\s*([A-H]+|正确|错误)\s*$", p.text.strip())
        if m:
            out.append(m.group(2))
    return out


def answer_file_of(paper_file: Path) -> Path:
    return paper_file.with_name(paper_file.name.replace("_", "_答案_", 1))



def _force_utf8() -> None:
    """Windows 控制台/CI 下强制 UTF-8 输出，避免中文打印触发 UnicodeEncodeError。"""
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")     # type: ignore[union-attr]
        except Exception:
            pass


_force_utf8()


def main() -> int:
    ap = argparse.ArgumentParser(description="回读生成的试卷做独立校验")
    ap.add_argument("--bank", required=True, help="题库文件")
    ap.add_argument("--out", required=True, help="试卷输出目录")
    ap.add_argument("--single", type=int, default=20)
    ap.add_argument("--multiple", type=int, default=10)
    ap.add_argument("--judge", type=int, default=10)
    ap.add_argument("--seed", type=int, default=12345)
    args = ap.parse_args()

    out_dir = Path(args.out)
    bank = load_bank(args.bank)
    by_stem = {norm(q.stem): q for q in bank.questions}
    plan = {"single": args.single, "multiple": args.multiple, "judge": args.judge}
    total_plan = sum(plan.values())

    print("=" * 78)
    print(f"题库：{args.bank}（{len(bank.questions)} 题）   输出：{out_dir}")
    print(f"蓝图：每卷 {total_plan} 题（{args.single}/{args.multiple}/{args.judge}）")
    print("=" * 78)

    papers = sorted((p for p in out_dir.glob("*.docx") if "答案" not in p.name), key=lambda p: p.name)
    expect = max(
        -(-len([q for q in bank.questions if q.qtype is QType.SINGLE]) // max(args.single, 1)),
        -(-len([q for q in bank.questions if q.qtype is QType.MULTIPLE]) // max(args.multiple, 1)),
        -(-len([q for q in bank.questions if q.qtype is QType.JUDGE]) // max(args.judge, 1)),
    )
    check(len(papers) == expect, f"试卷张数 == {expect}（实际 {len(papers)}）")

    seen_all: set[str] = set()
    for pf in papers:
        items = parse_paper(pf)
        af = answer_file_of(pf)
        answers = parse_answers(af) if af.exists() else []
        label = pf.stem
        stats = {"single": 0, "multiple": 0, "judge": 0}
        mismatch: list[str] = []
        for i, (item, ans) in enumerate(zip(items, answers, strict=False), start=1):
            q = by_stem.get(norm(strip_blank(item["stem"])))
            if q is None:
                mismatch.append(f"#{i} 题干无法定位到题库")
                continue
            seen_all.add(q.qid)
            if q.qtype is QType.JUDGE:
                stats["judge"] += 1
                if item["options"]:
                    mismatch.append(f"#{i} 判断题不应带选项")
                if ans != q.answer_display:
                    mismatch.append(f"#{i} 判断题答案不符 {ans} != {q.answer_display}")
            else:
                stats["single" if q.qtype is QType.SINGLE else "multiple"] += 1
                got = {t for lab, t in item["options"] if lab in ans}
                want = {q.options["ABCDEFGH".index(c)] for c in q.answer_letters}
                if len(item["options"]) != q.option_count:
                    mismatch.append(f"#{i} 选项数量不符")
                if got != want:
                    mismatch.append(f"#{i} 答案指向文本不符 {sorted(got)} != {sorted(want)}")
        print(f"\n[{label}]")
        check(len(items) == total_plan, f"{label} 题量 == {total_plan}（实际 {len(items)}）")
        check(len(answers) == len(items), f"{label} 答案条数 == 题目数（{len(answers)}/{len(items)}）")
        check(stats == plan, f"{label} 题型数量 == {args.single}/{args.multiple}/{args.judge}（实际 "
                             f"{stats['single']}/{stats['multiple']}/{stats['judge']}）")
        check(not mismatch, f"{label} 答案与卷面自洽（问题 {len(mismatch)} 项）")
        for m in mismatch[:5]:
            print("        -", m)

    print(f"\n[覆盖性] 去重后命中题库 {len(seen_all)} / {len(bank.questions)} 题")
    check(len(seen_all) == len(bank.questions), "一轮覆盖题库全部题目")

    print("\n[发布一致性] 与源码同种子生成结果比对")
    report = generate(bank.questions, GenOptions(blueprint=Blueprint(args.single, args.multiple, args.judge),
                                                 seed=args.seed))
    for pf, paper in zip(papers, report.papers, strict=False):
        items = parse_paper(pf)
        af = answer_file_of(pf)
        check([norm(strip_blank(it.question.stem)) for it in paper.all_items()]
              == [norm(strip_blank(it["stem"])) for it in items], f"{paper.name} 题目顺序一致")
        check([it.answer_display for it in paper.all_items()]
              == (parse_answers(af) if af.exists() else []), f"{paper.name} 答案序列一致")

    print("\n[排版要素]")
    if papers:
        text = "\n".join(p.text for p in Document(papers[0]).paragraphs)
        check("姓名" in text and "成绩" in text, "卷头含考生信息栏")
        check(all(k in text for k in ("单选题", "多选题", "判断题")), "含三大题型标题")
        check("满分" in text, "含满分说明")
        judge_items = [it for it in parse_paper(papers[0]) if not it["options"]]
        check(all("（" in i["stem"] for i in judge_items), "判断题带作答括号")

    print()
    if failures:
        print(f"❌ 校验未通过，共 {len(failures)} 项")
        return 1
    print("✅ 全部校验通过")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
