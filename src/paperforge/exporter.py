# -*- coding: utf-8 -*-
"""导出器：把生成的试卷渲染为 Word(.docx)、可打印 HTML、纯文本(.txt)。

三种格式共享同一套"文本渲染"逻辑（题干编号、选项排版、答案区），
Word 端额外做字体与版式设置（宋体 + 东亚字体，避免中文乱码/变宋体失败）。
"""
from __future__ import annotations

import html as html_lib
import re
from dataclasses import dataclass, field
from pathlib import Path

from .generator import GenReport
from .models import Paper, PaperItem, QType

__all__ = ["ExportOptions", "ExportResult", "export_all", "paper_to_text",
           "answer_to_text", "paper_to_html", "answer_to_html", "export_docx",
           "export_html", "export_txt", "safe_filename"]

_CN_NUM = "一二三四五六七八九十"


@dataclass
class ExportOptions:
    """导出参数。"""

    out_dir: Path = Path(".")
    formats: tuple[str, ...] = ("docx",)
    answers_separate: bool = True      # 答案单独成文件
    include_analysis: bool = False     # 答案中附带解析
    merge: bool = False                # 全部试卷合并为一个文件
    show_score: bool = True            # 显示分值
    scores: dict = field(default_factory=lambda: {"single": 1.0, "multiple": 2.0, "judge": 1.0})
    title: str = "试卷"
    bank_name: str = ""
    exam_minutes: int = 0              # 0 表示不显示考试时长
    base_name: str = ""

    def score_of(self, qtype: QType) -> float:
        return float(self.scores.get(qtype.key, 1.0))


@dataclass
class ExportResult:
    """导出结果。"""

    files: list[Path] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return bool(self.files) and not self.errors


# ---------------------------------------------------------------- 通用渲染
def safe_filename(name: str, max_len: int = 80) -> str:
    """清洗为 Windows 合法文件名。"""
    cleaned = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", name).strip(" .")
    cleaned = re.sub(r"_{2,}", "_", cleaned)
    return (cleaned or "未命名")[:max_len]


def _stem_with_blank(item: PaperItem) -> str:
    stem = item.stem.strip()
    if item.qtype is QType.JUDGE:
        tail = stem.rstrip()
        if not (tail.endswith("）") or tail.endswith(")")):
            return f"{tail}（　　）"
    return stem


def _section_title(order: int, qtype: QType, items: list[PaperItem], opts: ExportOptions) -> str:
    prefix = _CN_NUM[order] if order < len(_CN_NUM) else str(order + 1)
    n = len(items)
    head = f"{prefix}、{qtype.value}"
    if opts.show_score:
        per = opts.score_of(qtype)
        per_txt = f"{per:g}"
        return f"{head}（本大题共 {n} 题，每题 {per_txt} 分，共 {n * per:g} 分）"
    return f"{head}（本大题共 {n} 题）"


def _paper_header_lines(paper: Paper, opts: ExportOptions, index: int) -> list[str]:
    """卷头：标题 + 说明 + 考生信息。"""
    label = f"{opts.title} {paper.label}卷" if opts.title else f"{paper.label}卷"
    parts = [f"{len(paper.items_of(QType.SINGLE))} 道单选",
             f"{len(paper.items_of(QType.MULTIPLE))} 道多选",
             f"{len(paper.items_of(QType.JUDGE))} 道判断"]
    parts = [p for p in parts if not p.startswith("0 ")]
    desc = "本卷共 %d 题（%s）" % (paper.total, "、".join(parts))
    total_score = 0.0
    for qtype, items in paper.items.items():
        if opts.show_score:
            total_score += len(items) * opts.score_of(qtype)
    if opts.show_score:
        desc += f"，满分 {total_score:g} 分"
    if opts.exam_minutes:
        desc += f"，考试时间 {opts.exam_minutes} 分钟"
    info = "姓名：____________    部门：____________    工号：____________    成绩：__________"
    return [label, desc, info]


def paper_to_text(paper: Paper, opts: ExportOptions, index: int = 0,
                  include_answers: bool = False, include_analysis: bool = False) -> str:
    """试卷纯文本渲染（同时用于 TXT 导出与界面预览）。"""
    lines: list[str] = []
    head = _paper_header_lines(paper, opts, index)
    lines.append(head[0])
    lines.append(head[1])
    if head[2]:
        lines.append(head[2])
    lines.append("")
    order = 0
    no = 0
    for qtype in (QType.SINGLE, QType.MULTIPLE, QType.JUDGE):
        items = paper.items_of(qtype)
        if not items:
            continue
        lines.append(_section_title(order, qtype, items, opts))
        order += 1
        for item in items:
            no += 1
            lines.append(f"{no}. {_stem_with_blank(item)}")
            for label, text in zip(item.labels, item.texts, strict=False):
                lines.append(f"    {label}. {text}")
            lines.append("")
    if include_answers:
        lines.append("")
        lines.append(f"{paper.label}卷 参考答案")
        lines.append(answer_to_text(paper, opts, index, include_analysis=include_analysis))
    return "\n".join(lines).rstrip() + "\n"


def answer_to_text(paper: Paper, opts: ExportOptions, index: int = 0,
                   include_analysis: bool = False) -> str:
    """答案区纯文本渲染。"""
    lines: list[str] = [f"{opts.title} {paper.label}卷 参考答案" if opts.title else f"{paper.label}卷 参考答案", ""]
    for i, item in enumerate(paper.all_items(), start=1):
        text = f"{i}. {item.answer_display}"
        if include_analysis and item.question.analysis:
            text += f"    解析：{item.question.analysis}"
        lines.append(text)
    return "\n".join(lines).rstrip() + "\n"


# ---------------------------------------------------------------- TXT
def export_txt(report: GenReport, opts: ExportOptions) -> list[Path]:
    out: list[Path] = []
    out_dir = Path(opts.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    base = opts.base_name or "试卷"

    if opts.merge:
        buf: list[str] = []
        abuf: list[str] = []
        for i, paper in enumerate(report.papers):
            buf.append(paper_to_text(paper, opts, i, include_answers=not opts.answers_separate,
                                     include_analysis=opts.include_analysis))
            buf.append("\n" + "=" * 60 + "\n")
            if opts.answers_separate:
                abuf.append(answer_to_text(paper, opts, i, opts.include_analysis))
                abuf.append("\n" + "-" * 60 + "\n")
        f = out_dir / safe_filename(f"{base}_全部{report.paper_count}张.txt")
        f.write_text("\n".join(buf), encoding="utf-8-sig")
        out.append(f)
        if opts.answers_separate and abuf:
            fa = out_dir / safe_filename(f"{base}_答案_全部{report.paper_count}张.txt")
            fa.write_text("\n".join(abuf), encoding="utf-8-sig")
            out.append(fa)
    else:
        for i, paper in enumerate(report.papers):
            f = out_dir / safe_filename(f"{base}_{paper.label}卷.txt")
            f.write_text(paper_to_text(paper, opts, i, include_answers=not opts.answers_separate,
                                       include_analysis=opts.include_analysis), encoding="utf-8-sig")
            out.append(f)
            if opts.answers_separate:
                fa = out_dir / safe_filename(f"{base}_答案_{paper.label}卷.txt")
                fa.write_text(answer_to_text(paper, opts, i, opts.include_analysis), encoding="utf-8-sig")
                out.append(fa)
    return out


# ---------------------------------------------------------------- HTML
_HTML_CSS = """
@page { size: A4; margin: 18mm 16mm; }
body { font-family: "Microsoft YaHei", "SimSun", "Songti SC", sans-serif; font-size: 11pt;
       line-height: 1.75; color: #111; margin: 0; padding: 0 6mm; }
h1 { font-size: 17pt; text-align: center; margin: 0 0 4pt; }
p.desc { text-align: center; color: #333; margin: 0 0 4pt; }
p.info { margin: 0 0 10pt; letter-spacing: 0.5px; }
h2 { font-size: 13pt; margin: 14pt 0 6pt; border-left: 4px solid #333; padding-left: 8px; }
div.q { margin: 0 0 9pt; page-break-inside: avoid; }
div.stem { font-weight: 600; }
div.opt { padding-left: 2em; }
div.analysis { color: #555; font-size: 10pt; padding-left: 2em; }
section.paper { page-break-after: always; }
section.paper:last-child { page-break-after: auto; }
hr { border: none; border-top: 1px dashed #999; margin: 12pt 0; }
@media print { .noprint { display: none; } }
"""


def _html_escape(text: str) -> str:
    return html_lib.escape(text, quote=False)


def paper_to_html(paper: Paper, opts: ExportOptions, index: int = 0,
                  include_answers: bool = False) -> str:
    head = _paper_header_lines(paper, opts, index)
    parts = [f"<h1>{_html_escape(head[0])}</h1>", f'<p class="desc">{_html_escape(head[1])}</p>',
             f'<p class="info">{_html_escape(head[2])}</p>']
    order = 0
    no = 0
    for qtype in (QType.SINGLE, QType.MULTIPLE, QType.JUDGE):
        items = paper.items_of(qtype)
        if not items:
            continue
        parts.append(f"<h2>{_html_escape(_section_title(order, qtype, items, opts))}</h2>")
        order += 1
        for item in items:
            no += 1
            block = [f'<div class="q"><div class="stem">{no}. {_html_escape(_stem_with_blank(item))}</div>']
            for label, text in zip(item.labels, item.texts, strict=False):
                block.append(f'<div class="opt">{label}. {_html_escape(text)}</div>')
            block.append("</div>")
            parts.append("\n".join(block))
    if include_answers:
        parts.append("<hr>")
        parts.append(f"<h2>{_html_escape(paper.label)}卷 参考答案</h2>")
        for i, item in enumerate(paper.all_items(), start=1):
            line = f'<div class="q"><div class="stem">{i}. {_html_escape(item.answer_display)}</div>'
            if opts.include_analysis and item.question.analysis:
                line += f'<div class="analysis">解析：{_html_escape(item.question.analysis)}</div>'
            line += "</div>"
            parts.append(line)
    return "\n".join(parts)


def answer_to_html(paper: Paper, opts: ExportOptions, index: int = 0) -> str:
    parts = [f"<h1>{_html_escape(opts.title)} {paper.label}卷 参考答案</h1>"]
    for i, item in enumerate(paper.all_items(), start=1):
        line = f'<div class="q"><div class="stem">{i}. {_html_escape(item.answer_display)}</div>'
        if opts.include_analysis and item.question.analysis:
            line += f'<div class="analysis">解析：{_html_escape(item.question.analysis)}</div>'
        line += "</div>"
        parts.append(line)
    return "\n".join(parts)


def _html_doc(body: str, title: str) -> str:
    return (
        "<!DOCTYPE html>\n<html lang=\"zh-CN\">\n<head>\n"
        "<meta charset=\"utf-8\">\n"
        f"<title>{_html_escape(title)}</title>\n"
        f"<style>{_HTML_CSS}</style>\n</head>\n<body>\n{body}\n"
        "<p class=\"noprint\" style=\"text-align:center;color:#888;font-size:10pt;\">"
        "提示：浏览器中按 Ctrl+P 可打印或另存为 PDF</p>\n"
        "</body>\n</html>\n"
    )


def export_html(report: GenReport, opts: ExportOptions) -> list[Path]:
    out: list[Path] = []
    out_dir = Path(opts.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    base = opts.base_name or "试卷"

    if opts.merge:
        body = []
        for i, paper in enumerate(report.papers):
            body.append(f'<section class="paper">'
                        f'{paper_to_html(paper, opts, i, include_answers=not opts.answers_separate)}'
                        f'</section>')
        f = out_dir / safe_filename(f"{base}_全部{report.paper_count}张.html")
        f.write_text(_html_doc("\n".join(body), f"{opts.title} 全部{report.paper_count}张"), encoding="utf-8")
        out.append(f)
        if opts.answers_separate:
            ab = [f'<section class="paper">{answer_to_html(paper, opts, i)}</section>'
                  for i, paper in enumerate(report.papers)]
            fa = out_dir / safe_filename(f"{base}_答案_全部{report.paper_count}张.html")
            fa.write_text(_html_doc("\n".join(ab), f"{opts.title} 答案"), encoding="utf-8")
            out.append(fa)
    else:
        for i, paper in enumerate(report.papers):
            body = paper_to_html(paper, opts, i, include_answers=not opts.answers_separate)
            f = out_dir / safe_filename(f"{base}_{paper.label}卷.html")
            f.write_text(_html_doc(body, f"{opts.title} {paper.label}卷"), encoding="utf-8")
            out.append(f)
            if opts.answers_separate:
                fa = out_dir / safe_filename(f"{base}_答案_{paper.label}卷.html")
                fa.write_text(_html_doc(answer_to_html(paper, opts, i), f"{opts.title} {paper.label}卷 答案"),
                              encoding="utf-8")
                out.append(fa)
    return out


# ---------------------------------------------------------------- DOCX
def _set_font(run, size: float = 10.5, bold: bool = False, name: str = "宋体") -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        from docx.oxml import OxmlElement

        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)


def _add_paragraph(doc, text: str = "", *, size: float = 10.5, bold: bool = False,
                   align=None, indent: float = 0.0, space_after: float = 2.0,
                   color=None, first_line_indent: float = 0.0, name: str = "宋体"):
    from docx.shared import Cm, Pt, RGBColor

    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.25
    if indent:
        pf.left_indent = Cm(indent)
    if first_line_indent:
        pf.first_line_indent = Cm(first_line_indent)
    if text:
        run = p.add_run(text)
        _set_font(run, size, bold, name)
        if color is not None:
            run.font.color.rgb = RGBColor(*color)
    return p


def _setup_document():
    from docx import Document
    from docx.shared import Cm

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    from docx.oxml.ns import qn

    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        from docx.oxml import OxmlElement

        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), "宋体")
    rfonts.set(qn("w:hAnsi"), "宋体")
    rfonts.set(qn("w:eastAsia"), "宋体")
    return doc


def _write_paper_docx(doc, paper: Paper, opts: ExportOptions, index: int,
                      include_answers: bool) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    head = _paper_header_lines(paper, opts, index)
    _add_paragraph(doc, head[0], size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    _add_paragraph(doc, head[1], size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    _add_paragraph(doc, head[2], size=10.5, space_after=8)

    order = 0
    no = 0
    for qtype in (QType.SINGLE, QType.MULTIPLE, QType.JUDGE):
        items = paper.items_of(qtype)
        if not items:
            continue
        _add_paragraph(doc, _section_title(order, qtype, items, opts), size=12, bold=True,
                       space_after=4)
        order += 1
        for item in items:
            no += 1
            _add_paragraph(doc, f"{no}. {_stem_with_blank(item)}", size=10.5, space_after=1)
            longs = [len(t) for t in item.texts]
            if item.texts and max(longs) <= 12:
                joined = "　　".join(f"{lab}. {txt}" for lab, txt in zip(item.labels, item.texts, strict=False))
                _add_paragraph(doc, joined, size=10.5, indent=0.6, space_after=4)
            else:
                for lab, txt in zip(item.labels, item.texts, strict=False):
                    _add_paragraph(doc, f"{lab}. {txt}", size=10.5, indent=0.6, space_after=0)
                _add_paragraph(doc, "", size=6, space_after=2)

    if include_answers:
        doc.add_page_break()
        _write_answer_docx(doc, paper, opts)


def _write_answer_docx(doc, paper: Paper, opts: ExportOptions) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    _add_paragraph(doc, f"{opts.title} {paper.label}卷 参考答案", size=14, bold=True,
                   align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    for i, item in enumerate(paper.all_items(), start=1):
        _add_paragraph(doc, f"{i}. {item.answer_display}", size=10.5, space_after=1)
        if opts.include_analysis and item.question.analysis:
            _add_paragraph(doc, f"解析：{item.question.analysis}", size=9.5, indent=0.6,
                           space_after=4, color=(0x55, 0x55, 0x55))


def export_docx(report: GenReport, opts: ExportOptions) -> list[Path]:
    out: list[Path] = []
    out_dir = Path(opts.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    base = opts.base_name or "试卷"

    if opts.merge:
        doc = _setup_document()
        for i, paper in enumerate(report.papers):
            if i:
                doc.add_page_break()
            _write_paper_docx(doc, paper, opts, i, include_answers=False)
        f = out_dir / safe_filename(f"{base}_全部{report.paper_count}张.docx")
        doc.save(f)
        out.append(f)
        if opts.answers_separate:
            adoc = _setup_document()
            for i, paper in enumerate(report.papers):
                if i:
                    adoc.add_page_break()
                _write_answer_docx(adoc, paper, opts)
            fa = out_dir / safe_filename(f"{base}_答案_全部{report.paper_count}张.docx")
            adoc.save(fa)
            out.append(fa)
    else:
        for i, paper in enumerate(report.papers):
            if opts.answers_separate:
                doc = _setup_document()
                _write_paper_docx(doc, paper, opts, i, include_answers=False)
                f = out_dir / safe_filename(f"{base}_{paper.label}卷.docx")
                doc.save(f)
                out.append(f)
                adoc = _setup_document()
                _write_answer_docx(adoc, paper, opts)
                fa = out_dir / safe_filename(f"{base}_答案_{paper.label}卷.docx")
                adoc.save(fa)
                out.append(fa)
            else:
                doc = _setup_document()
                _write_paper_docx(doc, paper, opts, i, include_answers=True)
                f = out_dir / safe_filename(f"{base}_{paper.label}卷.docx")
                doc.save(f)
                out.append(f)
    return out


# ---------------------------------------------------------------- 统一入口
def export_all(report: GenReport, opts: ExportOptions) -> ExportResult:
    """按 ``opts.formats`` 导出全部格式，单个格式失败不影响其他格式。"""
    result = ExportResult()
    for fmt in opts.formats:
        fmt = fmt.lower().strip()
        try:
            if fmt == "docx":
                result.files.extend(export_docx(report, opts))
            elif fmt == "html":
                result.files.extend(export_html(report, opts))
            elif fmt == "txt":
                result.files.extend(export_txt(report, opts))
            else:
                result.errors.append(f"不支持的导出格式：{fmt}")
        except Exception as exc:  # 单个格式失败不影响其他格式
            result.errors.append(f"{fmt.upper()} 导出失败：{exc}")
    return result
